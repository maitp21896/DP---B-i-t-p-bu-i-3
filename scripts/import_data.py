"""Re-import the supplied workbook: python scripts/import_data.py (requires openpyxl)."""
import json, hashlib
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parent.parent
source=root/'data/source.xlsx'
w=openpyxl.load_workbook(source,data_only=False)
ids=['broad','lal','rt']
sets=[{'id':ids[i], 'name':w['Tong quan'].cell(i+4,1).value,'reach':w['Tong quan'].cell(i+4,3).value,'frequency':w['Tong quan'].cell(i+4,4).value} for i in range(3)]
lookup={s['name']:s['id'] for s in sets}
daily=[]
for r in list(w['Du lieu theo ngay'].values)[1:]:
 daily.append(dict(zip(['day','id','spend','clicks','lpv','atc','purchases'],[int(r[0][1:]),lookup[r[1]],*r[2:7]])))
assert len(daily)==21
assert all(sum(d['id']==s['id'] for d in daily)==7 for s in sets)
# Verify source aggregate formulas point to the daily source; cached formula results are absent.
for row in range(4,7):
 for col,src in [(2,"C"),(5,"D"),(6,"E"),(8,"F"),(10,"G")]:
  expected=f"=SUMIF('Du lieu theo ngay'!$B$2:$B$22,A{row},'Du lieu theo ngay'!${src}$2:${src}$22)"
  assert w['Tong quan'].cell(row,col).value==expected, "Unexpected aggregate formula: report discrepancy before import"
output={'source':'Glow_Nature_Campaign_Rescue_Formula_ATC_Updated.xlsx','sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'adSets':sets,'daily':daily,'limitations':['Ô công thức Excel không có kết quả lưu sẵn; tổng và tỷ lệ được tính lại từ 21 dòng dữ liệu gốc.','Reach và Frequency là số báo cáo được cung cấp, không cộng Reach theo ngày hoặc giữa các Ad Set.','Không có impressions, dữ liệu thiết bị/placement hoặc tốc độ trang; chưa tính được CTR/CPM hoặc kết luận nguyên nhân.']}
(root/'data/campaign.json').write_text(json.dumps(output,ensure_ascii=False,indent=2)+'\n')
print('Imported',len(daily),'daily rows; totals:',{k:sum(d[k] for d in daily) for k in ['spend','clicks','lpv','atc','purchases']})
